import sys
from datetime import datetime
from pathlib import Path
import requests
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication, QDialog, QLabel, QPushButton, QVBoxLayout
from openpyxl import load_workbook
from openpyxl.styles import Font
from copy import copy
import os


class InfoDialog(QDialog):
	"""通用信息弹窗，用于显示提示或错误信息"""

	def __init__(self, title: str, message: str, parent=None):
		super().__init__(parent)
		self.setWindowTitle(title)
		self.setModal(True)
		self.setFixedSize(350, 180)

		self.setStyleSheet("""
			QDialog {
				background-color: #ffffff;
			}
			QLabel {
				color: #2c3e50;
				font-size: 14px;
				font-family: "Segoe UI", "Microsoft YaHei", sans-serif;
				padding: 8px;
			}
			QPushButton {
				background-color: #ecf0f1;
				color: #2c3e50;
				border: 1px solid #bdc3c7;
				border-radius: 3px;
				padding: 5px 16px;
				font-size: 13px;
				min-width: 70px;
			}
			QPushButton:hover {
				background-color: #dfe6e9;
				border-color: #95a5a6;
			}
			QPushButton:pressed {
				background-color: #bdc3c7;
			}
		""")

		layout = QVBoxLayout(self)
		layout.setSpacing(15)
		layout.setContentsMargins(20, 20, 20, 20)

		label = QLabel(message, self)
		label.setAlignment(Qt.AlignmentFlag.AlignCenter)
		label.setWordWrap(True)

		ok_button = QPushButton("确定", self)
		ok_button.clicked.connect(self.accept)

		layout.addWidget(label)
		layout.addWidget(ok_button, alignment=Qt.AlignmentFlag.AlignCenter)


def show_info_dialog(title: str, message: str):
	"""显示信息弹窗（假设 QApplication 已由外部创建）"""
	dialog = InfoDialog(title, message)
	dialog.exec()


def get_download_path() -> Path:
	"""从 config.txt 读取保存目录，若失败则返回默认路径，并在该路径下自动创建 'YYYY年/MM月/' 子目录后返回"""
	config_file = Path("config.txt")
	default_path = Path("F:/淘金站施工请销点情况控制表/【4】《车站施工请销点控制表》/")

	try:
		if config_file.exists():
			path_str = config_file.read_text(encoding="utf-8").strip()
			if path_str:
				base_path = Path(path_str)
			else:
				base_path = default_path
		else:
			base_path = default_path
	except Exception as e:
		print(f"读取配置文件失败: {e}")
		base_path = default_path

	# 获取当前年月并创建子目录
	now = datetime.now()
	year_dir = f"{now.year}年"
	month_dir = f"{now.month}月"
	target_path = base_path / year_dir / month_dir
	target_path.mkdir(parents=True, exist_ok=True)

	return target_path


def download_night_construction_report() -> bytes | None:
	"""
	执行网络请求，下载「夜间施工台账」Excel 文件
	返回文件内容（字节）或 None（失败时）
	"""
	url = "http://lmis-shigong.subway.com/chartshow/sgyx/czsgyx.jsp"
	headers = {
		"Content-Type": "application/x-www-form-urlencoded",
		"User-Agent": (
			"Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
			"AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0.0.0 Safari/537.36"
		),
	}

	session = requests.Session()

	try:
		# 1. 打开页面（GET 请求，建立会话）
		session.get(url, headers=headers, timeout=10)

		# 2. 选择五号线
		session.post(url, headers=headers, data="line=LINE5&station=101&rq=", timeout=10)

		# 3. 点击「施工预想」（实际是提交一次表单，似乎不影响后续下载）
		session.post(
			url,
			headers=headers,
			data=f"line=LINE5&station=508&rq={datetime.now().strftime('%Y-%m-%d')}&b1=%E6%96%BD%E5%B7%A5%E9%A2%84%E6%83%B3",
			timeout=10,
		)

		# 4. 下载【夜间】施工台账
		today = datetime.now().strftime("%Y-%m-%d")
		download_data = (
			f"line=LINE5&station=508&rq={today}"
			"&b2=%E4%B8%8B%E8%BD%BD%E3%80%90%E5%A4%9C%E9%97%B4%E3%80%91%E6%96%BD%E5%B7%A5%E5%8F%B0%E8%B4%A6"
		)
		response = session.post(url, headers=headers, data=download_data, timeout=30)

		# 检查 HTTP 状态码
		response.raise_for_status()

		# 检查返回内容是否可能是 Excel 文件（简单判断 Content-Type 或文件头）
		content_type = response.headers.get("Content-Type", "")
		if "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" not in content_type:
			# 有些服务器可能返回 text/html 表示错误，我们检查前几个字节是否为 PK (Excel 文件头)
			if not response.content.startswith(b"PK"):
				raise ValueError("服务器返回的不是有效的 Excel 文件")

		return response.content

	except requests.exceptions.RequestException as e:
		show_info_dialog("网络错误", f"请求失败，请检查网络或站点访问权限。\n错误详情：{str(e)}")
		return None
	except Exception as e:
		show_info_dialog("数据错误", f"下载数据异常：{str(e)}")
		return None


def save_excel_file(content: bytes, save_dir: Path) -> bool:
	"""将字节内容保存为 Excel 文件，文件名格式：施工预想表MMDD.xlsx"""
	try:
		# 确保目录存在
		save_dir.mkdir(parents=True, exist_ok=True)

		today = datetime.now()
		filename = f"施工预想表{today.strftime('%m%d')}.xlsx"
		file_path = save_dir / filename

		with open(file_path, "wb") as f:
			f.write(content)

		# 使用 openpyxl 原地修改 A2 单元格，保留所有格式
		wb = load_workbook(file_path)
		ws = wb.active  # 获取当前活动工作表，或使用 wb['Sheet名']

		new_text = "次日__________，所有施工均已销点，所有列车已出清本站区域：行车值班员_______________，值班站长_______________"
		ws['A2'] = new_text  # 直接赋值，保留原有格式

		# 检测并修改 G8 和 G10 单元格
		set_font_size_if_long(ws['G8'])
		set_font_size_if_long(ws['G10'])

		wb.save(file_path)  # 保存覆盖原文件

		# 保存成功后自动打开 Excel 文件
		os.startfile(str(file_path))

		show_info_dialog("保存成功", f"文件已保存至：\n{file_path}")
		return True
	except Exception as e:
		show_info_dialog("保存失败", f"无法保存文件，请检查目录权限。\n错误详情：{str(e)}")
		return False

def set_font_size_if_long(cell, max_len=100, target_size=7.5):
	"""如果单元格内容长度超过 max_len，则将其字体大小设为 target_size，保留其他样式"""
	if cell.value is not None and len(str(cell.value)) > max_len:
		old_font = cell.font
		# 复制原有字体并修改大小
		new_font = Font(
			name=old_font.name,
			size=target_size,
			bold=old_font.bold,
			italic=old_font.italic,
			vertAlign=old_font.vertAlign,
			underline=old_font.underline,
			strike=old_font.strike,
			color=old_font.color,
			scheme=old_font.scheme,
			family=old_font.family,
			charset=old_font.charset,
			condense=old_font.condense,
			extend=old_font.extend,
			outline=old_font.outline,
			shadow=old_font.shadow
		)
		cell.font = new_font

def main():
	"""主流程：下载并保存夜间施工台账"""
	# 创建全局 QApplication（整个程序只需一个实例）
	app = QApplication(sys.argv)

	# 1. 下载文件
	file_content = download_night_construction_report()
	if file_content is None:
		# 下载失败，弹窗已在函数内部显示，直接退出
		sys.exit(1)

	# 2. 获取保存路径
	save_path = get_download_path()

	# 3. 保存文件
	save_excel_file(file_content, save_path)

	sys.exit(0)


if __name__ == "__main__":
	main()
